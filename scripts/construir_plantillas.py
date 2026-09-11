"""Construye las plantillas .docx (con tags {{ }} de docxtpl) a partir de la
estructura real observada en Docs/Oficios y Docs/Constancias. Se corre UNA
vez para generar los archivos en app/documents/templates_docx/ — esos
archivos sí se versionan en git (son plantillas de la app, no datos de
alumnos). Volver a correr este script solo si se necesita rediseñar una
plantilla desde cero.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent.parent
ASSETS = BASE_DIR / "app" / "documents" / "assets"
OUT = BASE_DIR / "app" / "documents" / "templates_docx"
OUT.mkdir(parents=True, exist_ok=True)
LOGO = ASSETS / "logo_udg.jpeg"


def _membrete(doc, lineas):
    header = doc.sections[0].header
    p_logo = header.paragraphs[0]
    p_logo.add_run().add_picture(str(LOGO), width=Cm(6))
    for texto in lineas:
        p = header.add_paragraph()
        p.add_run(texto)

    footer = doc.sections[0].footer
    for texto in [
        "Av. Universidad No. 203, Delegación Ixtapa, C.P. 48280.",
        "Puerto Vallarta, Jalisco. México.",
        "www.cuc.udg.mx",
    ]:
        fp = footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run(texto)
        run.font.size = Pt(8)


def _parrafo(doc, texto="", negrita=False, alineacion=None, tab_inicial=False):
    p = doc.add_paragraph()
    if alineacion is not None:
        p.alignment = alineacion
    run = p.add_run(("\t" if tab_inicial else "") + texto)
    run.bold = negrita
    return p


def construir_oficio_asignacion():
    doc = Document()
    _membrete(
        doc,
        [
            "CENTRO UNIVERSITARIO DE LA COSTA",
            "DIVISIÓN DE INGENIERIAS / DEPARTAMENTO DE CIENCIAS EXACTAS",
            "MAESTRÍA EN CIENCIAS EN GEOFÍSICA",
        ],
    )

    _parrafo(doc, "{{ oficio_numero }}")
    doc.add_paragraph()
    _parrafo(doc, "C. {{ alumno_nombre }}")
    _parrafo(doc, "Estudiante de la Maestría en Ciencias en Geofísica")
    _parrafo(doc, "P r e s e n t e")
    doc.add_paragraph()
    _parrafo(doc, "Asunto: Asignación de {{ rol_texto }} de tesis y autorización de trabajo de investigación.", negrita=True)
    doc.add_paragraph()
    _parrafo(
        doc,
        "Por este medio informo que la Junta Académica{% if acta_numero %} de acuerdo al Acta {{ acta_numero }}"
        "{% if acta_fecha %} de fecha {{ acta_fecha }}{% endif %}{% endif %}, se aprueba la designación de "
        "{{ profesor_nombre }} como su {{ rol_texto_corto }} para la dirección del "
        "trabajo de investigación{% if tesis_titulo %} titulado “{{ tesis_titulo }}”{% endif %}, de "
        "acuerdo al protocolo entregado.",
        tab_inicial=True,
    )
    doc.add_paragraph()
    _parrafo(doc, "Se pide tener en cuenta las siguientes obligaciones:")
    doc.add_paragraph()
    _parrafo(
        doc,
        "Es compromiso mutuo entre usted y su {{ rol_texto_corto }} entregar dos semanas antes de que se "
        "termine el semestre: 1) Avance de tesis semestral junto con su plan para el siguiente semestre, "
        "2) Presentación y defensa oral de los avances de la tesis, 3) Formato de evaluación del desempeño "
        "de becario (si aplica) y 4) Contestar los formatos de evaluación de profesores y materias del "
        "semestre.",
    )
    doc.add_paragraph()
    _parrafo(
        doc,
        "Durante el desarrollo de su investigación deberá respetar las normas de ética y evitar el plagio, "
        "procurando siempre las buenas prácticas de referencia y citado tanto en artículos como en cualquier "
        "publicación, en especial en la tesis en desarrollo.",
    )
    doc.add_paragraph()
    _parrafo(
        doc,
        "Es obligación {{ articulo_del_rol }} {{ rol_texto_corto }} dar seguimiento del desempeño académico y "
        "desarrollo de tesis del estudiante y reportar de manera concreta en los formatos correspondientes.",
    )
    doc.add_paragraph()
    _parrafo(doc, "Sin otro particular por el momento, quedo al pendiente de cualquier duda al respecto y envío un cordial saludo.")
    doc.add_paragraph()
    doc.add_paragraph()
    _parrafo(doc, "A t e n t a m e n t e", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "“PIENSA Y TRABAJA”", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "{% if lema_ciclo %}“{{ lema_ciclo }}”{% endif %}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "Puerto Vallarta, Jalisco, a {{ fecha_larga }}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    _parrafo(doc, "{{ coordinador_nombre }}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "Coordinador(a) de la Maestría en Geofísica", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    doc.add_paragraph()
    _parrafo(doc, "c.c.p. archivo")

    destino = OUT / "oficio_asignacion.docx"
    doc.save(destino)
    print("Generado:", destino)


def construir_constancia_direccion():
    doc = Document()
    _membrete(
        doc,
        [
            "CENTRO UNIVERSITARIO DE LA COSTA",
            "DIVISIÓN DE INGENIERIAS / DEPARTAMENTO DE CIENCIAS EXACTAS",
            "MAESTRÍA EN CIENCIAS EN GEOFÍSICA",
        ],
    )

    _parrafo(doc, "Constancia: {{ constancia_numero }}")
    doc.add_paragraph()
    _parrafo(doc, "A QUIEN CORRESPONDA")
    _parrafo(doc, "P r e s e n t e")
    doc.add_paragraph()
    _parrafo(
        doc,
        "El que subscribe, {{ coordinador_nombre }}, Coordinador de la Maestría en Ciencias en Geofísica del "
        "Centro Universitario de la Costa.",
    )
    doc.add_paragraph()
    _parrafo(
        doc,
        "Hace constar que {{ profesor_nombre }} fungió como {{ rol_texto }} de "
        "tesis del(la) alumno(a) {{ alumno_nombre }}{% if tesis_titulo %}, con el tema de tesis "
        "“{{ tesis_titulo }}”{% endif %}{% if fecha_defensa %} con fecha de defensa de "
        "{{ fecha_defensa }}{% endif %}.",
    )
    doc.add_paragraph()
    _parrafo(doc, "Se extiende la presente a petición del(la) interesado(a), para los usos y fines legales que a él/ella convengan.")
    doc.add_paragraph()
    doc.add_paragraph()
    _parrafo(doc, "ATENTAMENTE", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "“PIENSA Y TRABAJA”", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "{% if lema_ciclo %}“{{ lema_ciclo }}”{% endif %}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "Puerto Vallarta, Jal. {{ fecha_larga }}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    _parrafo(doc, "{{ coordinador_nombre }}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "Coordinador de la Maestría en Ciencias en Geofísica", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    doc.add_paragraph()
    _parrafo(doc, "c.c.p. Archivo")

    destino = OUT / "constancia_direccion.docx"
    doc.save(destino)
    print("Generado:", destino)


def construir_acta():
    doc = Document()
    header = doc.sections[0].header
    p_logo = header.paragraphs[0]
    p_logo.add_run().add_picture(str(LOGO), width=Cm(5))
    for texto in [
        "UNIVERSIDAD DE GUADALAJARA",
        "CENTRO UNIVERSITARIO DE LA COSTA",
        "SECRETARIA ACADEMICA / MAESTRÍA EN CIENCIAS EN GEOFÍSICA",
        "Acta {{ numero }}",
    ]:
        header.add_paragraph().add_run(texto)

    footer = doc.sections[0].footer
    for texto in [
        "Av. Universidad No. 203, Delegación Ixtapa, C.P. 48280",
        "Puerto Vallarta, Jalisco. México",
        "www.cuc.udg.mx",
    ]:
        fp = footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run(texto)
        run.font.size = Pt(8)

    _parrafo(doc, "Acta de Sesión de la Junta Académica Posgrado en Ciencias en Geofísica", negrita=True, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    _parrafo(
        doc,
        "En {{ lugar }}, siendo las {{ hora_inicio }} horas del día {{ fecha_larga }}"
        "{% if sede %}, en {{ sede }}{% endif %}, de conformidad al Artículo 14 Inciso I del Reglamento "
        "General de Posgrados se reunieron {{ asistentes }}, lo anterior a previa convocatoria que se les "
        "hizo para efectuar esta reunión que tiene como finalidad los siguientes puntos del orden del día:",
    )
    doc.add_paragraph()
    _parrafo(doc, "ORDEN DEL DIA:", negrita=True)
    _parrafo(doc, "{%for punto in puntos %}")
    _parrafo(doc, "{{ loop.index }}. {{ punto.titulo }}")
    _parrafo(doc, "{%endfor %}")
    doc.add_paragraph()
    _parrafo(doc, "RESOLUTIVOS", negrita=True)
    _parrafo(doc, "{%for punto in puntos %}")
    _parrafo(doc, "Como punto número {{ punto.numero_texto }} de la orden del día, {{ punto.resolutivo }}", tab_inicial=True)
    doc.add_paragraph()
    _parrafo(doc, "{%endfor %}")
    _parrafo(
        doc,
        "Y siendo las {{ hora_fin }} horas del día {{ fecha_larga }}, de conformidad con el último punto de "
        "la orden del día se da por clausurada la reunión de la Junta Académica del Posgrado en Ciencias en "
        "Geofísica, firmando la presente acta los que en ella intervinieron.",
    )
    doc.add_paragraph()
    doc.add_paragraph()
    _parrafo(doc, "ATENTAMENTE", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "“PIENSA Y TRABAJA”", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "{% if lema_ciclo %}“{{ lema_ciclo }}”{% endif %}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "Puerto Vallarta, Jal. {{ fecha_larga }}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    _parrafo(doc, "{{ coordinador_nombre }}", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, "Coordinador(a) de la Maestría en Ciencias en Geofísica", alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    doc.add_paragraph()
    _parrafo(doc, "(Espacio para firmas de los asistentes)")

    destino = OUT / "acta.docx"
    doc.save(destino)
    print("Generado:", destino)


if __name__ == "__main__":
    construir_oficio_asignacion()
    construir_constancia_direccion()
    construir_acta()
