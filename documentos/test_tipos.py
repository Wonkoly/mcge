import docx
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from documentos import tipos
from documentos.models import TipoDocumentoPersonalizado


def _docx_bytes(parrafos: list[str]) -> bytes:
    import io

    doc = docx.Document()
    for texto in parrafos:
        doc.add_paragraph(texto)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _archivo_docx(nombre: str, parrafos: list[str]) -> SimpleUploadedFile:
    return SimpleUploadedFile(
        nombre, _docx_bytes(parrafos),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def _texto_completo(buffer) -> str:
    doc = docx.Document(buffer)
    return "\n".join(p.text for p in doc.paragraphs)


@pytest.mark.django_db
def test_crear_tipo_dedupe_clave(settings, tmp_path):
    settings.PLANTILLAS_DIR = tmp_path
    uno = tipos.crear_tipo(etiqueta="Cambio de Estatus", categoria="oficio", archivo=_archivo_docx("a.docx", ["hola"]))
    dos = tipos.crear_tipo(etiqueta="Cambio de Estatus", categoria="oficio", archivo=_archivo_docx("b.docx", ["hola"]))
    assert uno.clave == "cambio_de_estatus"
    assert dos.clave == "cambio_de_estatus_2"


@pytest.mark.django_db
def test_crear_tipo_categoria_invalida_lanza_valueerror(settings, tmp_path):
    settings.PLANTILLAS_DIR = tmp_path
    with pytest.raises(ValueError):
        tipos.crear_tipo(etiqueta="X", categoria="no-existe", archivo=_archivo_docx("a.docx", ["x"]))


@pytest.mark.django_db
def test_crear_tipo_sin_archivo_docx_lanza_valueerror(settings, tmp_path):
    settings.PLANTILLAS_DIR = tmp_path
    no_docx = SimpleUploadedFile("a.txt", b"no es un docx", content_type="text/plain")
    with pytest.raises(ValueError):
        tipos.crear_tipo(etiqueta="X", categoria="oficio", archivo=no_docx)


@pytest.mark.django_db
def test_crear_tipo_queda_confirmado_de_inmediato(settings, tmp_path):
    settings.PLANTILLAS_DIR = tmp_path
    tipo = tipos.crear_tipo(etiqueta="Visita de Prueba", categoria="oficio", archivo=_archivo_docx("a.docx", ["{{ coordinador_nombre }}"]))
    assert tipo.estado == "confirmado"
    assert tipo.plantilla_archivo == "visita_de_prueba.docx"
    assert (settings.PLANTILLAS_DIR / tipo.plantilla_archivo).exists()


@pytest.mark.django_db
def test_reemplazar_plantilla_sobreescribe_el_mismo_archivo(settings, tmp_path):
    settings.PLANTILLAS_DIR = tmp_path
    tipo = tipos.crear_tipo(etiqueta="Prueba", categoria="oficio", archivo=_archivo_docx("a.docx", ["viejo"]))
    nombre_original = tipo.plantilla_archivo

    tipos.reemplazar_plantilla(tipo, _archivo_docx("b.docx", ["{{ coordinador_nombre }}"]))
    tipo.refresh_from_db()

    assert tipo.plantilla_archivo == nombre_original
    ruta = settings.PLANTILLAS_DIR / tipo.plantilla_archivo
    assert "coordinador_nombre" in "\n".join(p.text for p in docx.Document(str(ruta)).paragraphs)


@pytest.mark.django_db
def test_generar_vista_previa_sin_plantilla_lanza_error(settings, tmp_path):
    settings.PLANTILLAS_DIR = tmp_path
    tipo = TipoDocumentoPersonalizado.objects.create(clave="sin_archivo", etiqueta="Sin Archivo", categoria="oficio")
    with pytest.raises(tipos.PlantillaFaltanteError):
        tipos.generar_vista_previa(tipo)


@pytest.mark.django_db
def test_generar_vista_previa_sustituye_variables_de_ejemplo(settings, tmp_path):
    settings.PLANTILLAS_DIR = tmp_path
    tipo = tipos.crear_tipo(
        etiqueta="Vista Previa", categoria="oficio",
        archivo=_archivo_docx("a.docx", ["Firma: {{ coordinador_nombre }}"]),
    )

    buffer = tipos.generar_vista_previa(tipo)

    texto = _texto_completo(buffer)
    assert "{{" not in texto
    assert "Coordinador" in texto or "Rendón" in texto


@pytest.mark.django_db
def test_eliminar_tipo_borra_el_archivo(settings, tmp_path):
    settings.PLANTILLAS_DIR = tmp_path
    tipo = tipos.crear_tipo(etiqueta="Para Borrar", categoria="oficio", archivo=_archivo_docx("a.docx", ["x"]))
    ruta = settings.PLANTILLAS_DIR / tipo.plantilla_archivo
    assert ruta.exists()

    tipos.eliminar_tipo(tipo)

    assert not ruta.exists()
    assert not TipoDocumentoPersonalizado.objects.filter(pk=tipo.pk).exists()
